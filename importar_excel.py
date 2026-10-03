# importar_excel.py
# Quantum Core - Carga masiva desde Excel hacia quantum_wallet.db.     [AGREGADO]
#
# Lee carga_masiva_quantum.xlsx (hojas "usuarios" y "movimientos") e inserta cada
# fila con las mismas funciones de configurar_db.py, de modo que las llaves,
# las restricciones CHECK/UNIQUE y los triggers validan cada registro.
#
#   - Cada fila es atomica: se importa completa o no se importa.
#   - Las filas rechazadas quedan en errores_importacion.log y en
#     resultado_importacion.xlsx, con la hoja, la fila y el motivo.
#   - La tabla importaciones registra cada carga y evita importar dos veces
#     el mismo archivo (se compara su huella SHA-256).
#
# Requiere:  pip3 install openpyxl
# Orden:     python3 configurar_db.py  ->  python3 generar_excel.py  ->  python3 importar_excel.py

import hashlib
import logging
import sqlite3
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from configurar_db import (RUTA_DB, conciliar, conectar, realizar_pago, recargar,
                           registrar_empleado, registrar_usuario, transferir,
                           vincular_empleado)

CARPETA = Path(__file__).parent
RUTA_EXCEL = CARPETA / "carga_masiva_quantum.xlsx"
RUTA_LOG = CARPETA / "errores_importacion.log"
RUTA_RESULTADO = CARPETA / "resultado_importacion.xlsx"

# El log conserva el historial de todas las cargas (modo de agregar)
logging.basicConfig(
    filename=RUTA_LOG, filemode="a", level=logging.INFO, encoding="utf-8",
    format="%(asctime)s | %(levelname)-7s | %(message)s", datefmt="%Y-%m-%d %H:%M:%S",
)
log = logging.getLogger("importacion")

# Traduce el mensaje tecnico de SQLite a un motivo legible
MOTIVOS = [
    ("ck_wallet_saldo_no_negativo", "Saldo insuficiente"),
    ("ck_transaccion_monto_positivo", "Monto negativo o cero"),
    ("ck_transaccion_saldo", "Saldo insuficiente"),
    ("limite por transaccion", "Monto supera el limite"),
    ("ck_transaccion_contraparte", "Transferencia a la misma wallet"),
    ("usuarios.email", "Correo duplicado"),
    ("usuarios.cedula", "Cedula duplicada"),
    ("usuarios.nit", "NIT duplicado"),
    ("ck_usuario_nit_empresa", "Empresa sin NIT"),
    ("ck_usuario_cedula_empleado", "Empleado sin cedula"),
    ("documento no registrado", "Documento no registrado"),
    ("empresa no registrada", "Empresa no registrada"),
    ("tipo no reconocido", "Tipo de movimiento no valido"),
    ("fecha no valida", "Fecha no valida"),
]


def motivo(error):
    texto = str(error)
    for clave, etiqueta in MOTIVOS:
        if clave in texto:
            return etiqueta
    return "Otro"


# ----------------------------------------------------------------------
# Tabla de control de cargas
# ----------------------------------------------------------------------
def preparar_control(conexion):
    conexion.execute("""
        CREATE TABLE IF NOT EXISTS importaciones (
            id_importacion  INTEGER PRIMARY KEY AUTOINCREMENT,
            archivo         TEXT    NOT NULL,
            huella_sha256   TEXT    NOT NULL UNIQUE,
            fecha           TEXT    NOT NULL DEFAULT (datetime('now', 'localtime')),
            filas_leidas    INTEGER NOT NULL,
            importadas      INTEGER NOT NULL,
            rechazadas      INTEGER NOT NULL
        );
    """)
    conexion.commit()


def huella(ruta):
    return hashlib.sha256(ruta.read_bytes()).hexdigest()


# ----------------------------------------------------------------------
# Conversion de valores leidos del Excel
# ----------------------------------------------------------------------
def texto(valor):
    if valor is None or str(valor).strip() == "":
        return None
    return str(valor).strip()


def convertir_fecha(valor):
    if isinstance(valor, datetime):
        return valor.strftime("%Y-%m-%d %H:%M:%S")
    try:
        return datetime.strptime(str(valor).strip(), "%d/%m/%Y %H:%M").strftime("%Y-%m-%d %H:%M:%S")
    except ValueError:
        raise ValueError(f"fecha no valida: {valor}") from None


def id_por_documento(conexion, documento):
    fila = conexion.execute(
        "SELECT id_usuario FROM usuarios WHERE cedula = ? OR nit = ?;", (documento, documento)
    ).fetchone()
    if fila is None:
        raise ValueError(f"documento no registrado: {documento}")
    return fila[0]


# ----------------------------------------------------------------------
# Importacion por hoja
# ----------------------------------------------------------------------
def importar_usuarios(conexion, hoja, rechazos):
    leidas = importadas = 0
    filas = list(enumerate(hoja.iter_rows(min_row=2, values_only=True), start=2))
    # Orden de dependencias: los empleados se cargan despues de sus empresas.
    # sorted() es estable, asi que dentro de cada grupo se respeta el orden del archivo.
    filas.sort(key=lambda par: texto(par[1][0]) == "EMPLEADO")
    for n, fila in filas:
        leidas += 1
        tipo, nombre, email, cedula, nit, cargo, salario, nit_empresa = fila
        try:
            tipo = texto(tipo)
            if tipo == "EMPLEADO":
                # Se valida la empresa ANTES de insertar: la fila debe quedar completa o no quedar
                fila_empresa = conexion.execute(
                    "SELECT id_usuario FROM usuarios WHERE nit = ? AND tipo_usuario = 'EMPRESA';",
                    (texto(nit_empresa),),
                ).fetchone()
                if fila_empresa is None:
                    raise ValueError(f"empresa no registrada: {nit_empresa}")
                id_empleado = registrar_empleado(conexion, texto(nombre), texto(cedula),
                                                 texto(cargo), float(salario), email=texto(email))
                vincular_empleado(conexion, fila_empresa[0], id_empleado)
            else:
                registrar_usuario(conexion, texto(nombre), email=texto(email),
                                  cedula=texto(cedula), nit=texto(nit), tipo=tipo)
            importadas += 1
        except (sqlite3.DatabaseError, ValueError, TypeError) as error:
            conexion.rollback()
            rechazos.append(("usuarios", n, f"{tipo} | {nombre}", motivo(error), str(error)))
            log.warning("usuarios fila %d | %s | %s -> %s", n, tipo, nombre, error)
    return leidas, importadas


def importar_movimientos(conexion, hoja, rechazos):
    leidas = importadas = 0
    for n, fila in enumerate(hoja.iter_rows(min_row=2, values_only=True), start=2):
        leidas += 1
        fecha, tipo, origen, destino, monto, descripcion = fila
        try:
            fecha_sql = convertir_fecha(fecha)
            monto = float(monto)
            id_origen = id_por_documento(conexion, texto(origen))
            if tipo == "RECARGA":
                recargar(conexion, id_origen, monto, descripcion, fecha=fecha_sql)
            elif tipo == "PAGO":
                realizar_pago(conexion, id_origen, monto, descripcion, fecha=fecha_sql)
            elif tipo == "TRANSFERENCIA":
                id_destino = id_por_documento(conexion, texto(destino))
                transferir(conexion, id_origen, id_destino, monto, descripcion, fecha=fecha_sql)
            else:
                raise ValueError(f"tipo no reconocido: {tipo}")
            importadas += 1
        except (sqlite3.DatabaseError, ValueError, TypeError) as error:
            conexion.rollback()
            rechazos.append(("movimientos", n, f"{tipo} | {origen} | {monto}", motivo(error), str(error)))
            log.warning("movimientos fila %d | %s | %s | %s -> %s", n, tipo, origen, monto, error)
    return leidas, importadas


def exportar_rechazos(rechazos):
    libro = Workbook()
    hoja = libro.active
    hoja.title = "rechazados"
    hoja.append(["hoja", "fila", "registro", "motivo", "detalle tecnico"])
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="15304A")
    for rechazo in rechazos:
        hoja.append(list(rechazo))
    for letra, ancho in zip("ABCDE", (14, 7, 48, 32, 60)):
        hoja.column_dimensions[letra].width = ancho
    hoja.freeze_panes = "A2"
    libro.save(RUTA_RESULTADO)


# ----------------------------------------------------------------------
# Programa principal
# ----------------------------------------------------------------------
if __name__ == "__main__":
    print("=" * 66)
    print(" QUANTUM CORE - Carga masiva desde Excel")
    print("=" * 66)
    if not RUTA_DB.exists():
        raise SystemExit("No existe quantum_wallet.db: ejecute primero configurar_db.py")
    if not RUTA_EXCEL.exists():
        raise SystemExit("No existe el Excel: ejecute primero generar_excel.py")

    conexion = conectar()
    preparar_control(conexion)
    firma = huella(RUTA_EXCEL)
    previa = conexion.execute(
        "SELECT id_importacion, fecha FROM importaciones WHERE huella_sha256 = ?;", (firma,)
    ).fetchone()
    if previa:
        print(f"Archivo: {RUTA_EXCEL.name}")
        print(f"CARGA RECHAZADA: este archivo ya se importo (importacion #{previa[0]}, {previa[1]}).")
        print("Para repetir la prueba, ejecute configurar_db.py y vuelva a importar.")
        log.error("Archivo ya importado (importacion #%d): %s", previa[0], RUTA_EXCEL.name)
        raise SystemExit(1)

    libro = load_workbook(RUTA_EXCEL, read_only=True, data_only=True)
    trans_antes = conexion.execute("SELECT count(*) FROM transacciones;").fetchone()[0]
    log.info("Inicio de la importacion: %s", RUTA_EXCEL.name)

    rechazos = []
    resumen = {}
    for nombre_hoja, funcion in (("usuarios", importar_usuarios),
                                 ("movimientos", importar_movimientos)):
        leidas, importadas = funcion(conexion, libro[nombre_hoja], rechazos)
        resumen[nombre_hoja] = (leidas, importadas, leidas - importadas)
    libro.close()

    total_leidas = sum(r[0] for r in resumen.values())
    total_importadas = sum(r[1] for r in resumen.values())
    with conexion:
        conexion.execute(
            "INSERT INTO importaciones (archivo, huella_sha256, filas_leidas, importadas, rechazadas) "
            "VALUES (?, ?, ?, ?, ?);",
            (RUTA_EXCEL.name, firma, total_leidas, total_importadas, len(rechazos)),
        )
    exportar_rechazos(rechazos)

    print(f"Archivo: {RUTA_EXCEL.name}\n")
    print(f"    {'Hoja':<13}{'Leidas':>8}{'Importadas':>12}{'Rechazadas':>12}")
    for nombre_hoja, (leidas, importadas, rechazadas) in resumen.items():
        print(f"    {nombre_hoja:<13}{leidas:>8}{importadas:>12}{rechazadas:>12}")
    print(f"    {'TOTAL':<13}{total_leidas:>8}{total_importadas:>12}{len(rechazos):>12}")

    print("\n    Rechazos por motivo:")
    for etiqueta, cantidad in Counter(r[3] for r in rechazos).most_common():
        print(f"      {cantidad:>3}  {etiqueta}")

    usuarios = conexion.execute("SELECT count(*) FROM usuarios;").fetchone()[0]
    trans = conexion.execute("SELECT count(*) FROM transacciones;").fetchone()[0]
    print("\n    Estado de la base de datos:")
    print(f"      Usuarios registrados      : {usuarios}")
    print(f"      Transacciones registradas : {trans} ({trans - trans_antes} nuevas)")
    descuadres = conciliar(conexion)
    print("      Conciliacion de saldos    :",
          "correcta en todas las wallets" if not descuadres else f"descuadre en {descuadres}")
    print(f"\n    Detalle de rechazos: {RUTA_LOG.name} y {RUTA_RESULTADO.name}")

    log.info("Fin: %d filas leidas, %d importadas, %d rechazadas",
             total_leidas, total_importadas, len(rechazos))
    conexion.close()
