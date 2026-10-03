# generar_excel.py
# Quantum Core - Genera el archivo carga_masiva_quantum.xlsx con datos simulados
# para la carga masiva de la base de datos quantum_wallet.db.        [AGREGADO]
#
#   Hoja "usuarios"    : 60 filas  (personas, empresas y empleados)
#   Hoja "movimientos" : 500 filas (recargas, pagos y transferencias de septiembre de 2026)
#
# Aproximadamente el 5 % de las filas contiene errores intencionales para comprobar
# que la importacion rechaza los registros invalidos y los deja en el log.
# La semilla fija hace que el archivo sea identico en cada ejecucion.
#
# Requiere:  pip3 install openpyxl
# Ejecutar:  python3 generar_excel.py

import random
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RUTA_EXCEL = Path(__file__).with_name("carga_masiva_quantum.xlsx")
SEMILLA = 2026
TOTAL_MOVIMIENTOS = 500
LIMITE_TRANSACCION = 50_000_000

NOMBRES = ["Andres", "Camila", "Santiago", "Valentina", "Juan Pablo", "Mariana", "Sebastian",
           "Daniela", "Alejandro", "Laura", "Mateo", "Isabella", "Nicolas", "Sofia", "Felipe",
           "Manuela", "Carlos", "Paula", "Julian", "Natalia", "David", "Gabriela", "Esteban",
           "Catalina", "Diego", "Juliana", "Tomas", "Sara", "Miguel", "Luisa", "Jorge", "Ana Maria"]
APELLIDOS = ["Restrepo", "Gomez", "Ramirez", "Zuluaga", "Henao", "Ospina", "Giraldo", "Arango",
             "Echeverri", "Montoya", "Castano", "Velez", "Mejia", "Botero", "Jaramillo",
             "Cardona", "Rios", "Alzate", "Quintero", "Salazar", "Lopera", "Valencia"]
EMPRESAS = ["Flores del Oriente SAS", "Cafe La Montana SAS", "Logistica Andina SAS",
            "TecnoRetiro SAS", "Lacteos El Carmen SAS", "Ceramicas Viboral SAS",
            "Transportes La Ceja SAS", "Agroexport Rionegro SAS", "Panaderia San Antonio SAS",
            "Consultores Llanogrande SAS"]
CARGOS = [("Auxiliar administrativo", 1_750_000), ("Analista", 3_200_000),
          ("Desarrollador", 4_500_000), ("Coordinador", 5_800_000), ("Operario", 1_550_000),
          ("Contador", 4_100_000), ("Vendedor", 2_300_000)]
COMERCIOS = ["Pago en supermercado", "Pago de servicios publicos", "Pago en restaurante",
             "Pago de transporte", "Pago en farmacia", "Pago de plan de celular",
             "Pago en tienda de ropa", "Pago de matricula", "Pago en ferreteria"]


def sin_tildes(texto):
    return texto.lower().replace(" ", ".")


def generar_usuarios(rnd):
    """Devuelve las filas de la hoja usuarios y la lista de usuarios validos."""
    filas, validos = [], []
    cedulas = set()

    def nueva_cedula():
        while True:
            c = str(rnd.randint(1_000_000_000, 1_099_999_999))
            if c not in cedulas:
                cedulas.add(c)
                return c

    usados = set()

    def nuevo_nombre():
        while True:
            n = f"{rnd.choice(NOMBRES)} {rnd.choice(APELLIDOS)} {rnd.choice(APELLIDOS)}"
            if n not in usados:
                usados.add(n)
                return n

    # 9 empresas validas
    empresas = []
    for i, razon in enumerate(EMPRESAS[:9]):
        nit = f"9{rnd.randint(10_000_000, 99_999_999)}-{rnd.randint(0, 9)}"
        email = "pagos@" + razon.lower().replace(" sas", "").replace(" ", "") + ".co"
        fila = {"tipo_usuario": "EMPRESA", "nombre": razon, "email": email, "cedula": None,
                "nit": nit, "cargo": None, "salario": None, "nit_empresa": None}
        filas.append(fila); validos.append(fila); empresas.append(fila)

    # 33 personas validas
    for _ in range(33):
        nombre = nuevo_nombre()
        partes = nombre.split()
        email = f"{sin_tildes(partes[0])}.{sin_tildes(partes[-2])}{rnd.randint(1, 99)}@correo.com"
        fila = {"tipo_usuario": "PERSONA", "nombre": nombre, "email": email,
                "cedula": nueva_cedula(), "nit": None, "cargo": None, "salario": None,
                "nit_empresa": None}
        filas.append(fila); validos.append(fila)

    # 14 empleados validos, vinculados a una empresa
    for _ in range(14):
        nombre = nuevo_nombre()
        cargo, salario = rnd.choice(CARGOS)
        empresa = rnd.choice(empresas)
        fila = {"tipo_usuario": "EMPLEADO", "nombre": nombre, "email": None,
                "cedula": nueva_cedula(), "nit": None, "cargo": cargo, "salario": salario,
                "nit_empresa": empresa["nit"]}
        filas.append(fila); validos.append(fila)

    rnd.shuffle(filas)

    # 4 filas con errores intencionales
    persona = next(f for f in validos if f["tipo_usuario"] == "PERSONA")
    errores = [
        {"tipo_usuario": "PERSONA", "nombre": "Registro Duplicado", "email": persona["email"],
         "cedula": nueva_cedula(), "nit": None, "cargo": None, "salario": None, "nit_empresa": None},
        {"tipo_usuario": "EMPRESA", "nombre": EMPRESAS[9], "email": "info@consultores.co",
         "cedula": None, "nit": None, "cargo": None, "salario": None, "nit_empresa": None},
        {"tipo_usuario": "EMPLEADO", "nombre": "Empleado Sin Cedula", "email": None,
         "cedula": None, "nit": None, "cargo": "Operario", "salario": 1_550_000,
         "nit_empresa": empresas[0]["nit"]},
        {"tipo_usuario": "EMPLEADO", "nombre": "Empleado Sin Empresa", "email": None,
         "cedula": nueva_cedula(), "nit": None, "cargo": "Analista", "salario": 3_200_000,
         "nit_empresa": "999999999-9"},
    ]
    for error in errores:
        filas.insert(rnd.randint(5, len(filas)), error)
    return filas, validos


def documento(usuario):
    return usuario["nit"] or usuario["cedula"]


def generar_movimientos(rnd, validos):
    """Simula el mes de septiembre llevando el saldo de cada usuario."""
    saldo = {documento(u): 0.0 for u in validos}
    inicio = datetime(2026, 9, 1, 8, 0)
    movimientos = []

    def agregar(fecha, tipo, origen, destino, monto, descripcion):
        movimientos.append({"fecha": fecha, "tipo": tipo, "documento_origen": origen,
                            "documento_destino": destino, "monto": monto,
                            "descripcion": descripcion})

    # 1) Recarga inicial de cada usuario (1 de septiembre)
    for i, u in enumerate(validos):
        doc = documento(u)
        if u["tipo_usuario"] == "EMPRESA":
            monto = rnd.randrange(25_000_000, 45_000_000, 500_000)
        elif u["tipo_usuario"] == "EMPLEADO":
            monto = rnd.randrange(300_000, 1_500_000, 50_000)
        else:
            monto = rnd.randrange(800_000, 5_000_000, 50_000)
        saldo[doc] += monto
        agregar(inicio + timedelta(minutes=7 * i), "RECARGA", doc, None, float(monto),
                "Recarga inicial")

    # 2) Pago de nomina el 30 de septiembre (empresa -> empleado)
    nomina = []
    for u in validos:
        if u["tipo_usuario"] == "EMPLEADO":
            nomina.append((u["nit_empresa"], documento(u), float(u["salario"])))

    errores_previstos = 25
    aleatorios = TOTAL_MOVIMIENTOS - len(movimientos) - len(nomina) - errores_previstos

    # 3) Movimientos aleatorios entre el 2 y el 29 de septiembre
    fechas = sorted(datetime(2026, 9, 2, 7, 0) + timedelta(minutes=rnd.randint(0, 27 * 24 * 60))
                    for _ in range(aleatorios + errores_previstos))
    posiciones_error = set(rnd.sample(range(len(fechas)), errores_previstos))
    tipos_error = (["saldo"] * 6 + ["negativo"] * 4 + ["documento"] * 5 + ["limite"] * 3
                   + ["mismo"] * 3 + ["tipo"] * 2 + ["fecha"] * 2)
    rnd.shuffle(tipos_error)
    personas = [documento(u) for u in validos if u["tipo_usuario"] != "EMPRESA"]
    empresas = [documento(u) for u in validos if u["tipo_usuario"] == "EMPRESA"]
    todos = personas + empresas

    for k, fecha in enumerate(fechas):
        if k in posiciones_error:
            error = tipos_error.pop()
            doc = rnd.choice(personas)
            if error == "saldo":
                agregar(fecha, "PAGO", doc, None, round(saldo[doc] + rnd.randint(1, 3) * 1_000_000, 2),
                        "Pago en concesionario")
            elif error == "negativo":
                agregar(fecha, rnd.choice(["RECARGA", "PAGO"]), doc, None,
                        -float(rnd.randrange(10_000, 200_000, 1_000)), "Ajuste manual")
            elif error == "documento":
                agregar(fecha, "TRANSFERENCIA", doc, str(rnd.randint(70_000_000, 79_999_999)),
                        float(rnd.randrange(20_000, 300_000, 1_000)), "Transferencia a tercero")
            elif error == "limite":
                agregar(fecha, "RECARGA", rnd.choice(empresas), None,
                        float(rnd.randrange(60_000_000, 90_000_000, 1_000_000)), "Recarga por consignacion")
            elif error == "mismo":
                agregar(fecha, "TRANSFERENCIA", doc, doc, float(rnd.randrange(10_000, 100_000, 1_000)),
                        "Transferencia propia")
            elif error == "tipo":
                agregar(fecha, "RETIRO", doc, None, float(rnd.randrange(50_000, 300_000, 10_000)),
                        "Retiro en cajero")
            else:  # fecha inexistente: 31 de septiembre
                agregar(f"31/09/2026 {fecha:%H:%M}", "PAGO", doc, None,
                        float(rnd.randrange(10_000, 90_000, 1_000)), rnd.choice(COMERCIOS))
            continue

        opcion = rnd.random()
        if opcion < 0.45:                      # pago de una persona en un comercio
            doc = rnd.choice(personas)
            tope = min(saldo[doc] * 0.25, 600_000)
            if tope < 5_000:
                opcion = 0.9
            else:
                monto = round(rnd.uniform(5_000, tope), 2) if rnd.random() < 0.4 else \
                    float(rnd.randrange(5_000, int(tope), 1_000) if tope > 6_000 else 5_000)
                saldo[doc] -= monto
                agregar(fecha, "PAGO", doc, None, monto, rnd.choice(COMERCIOS))
                continue
        if opcion < 0.60:                      # pago de una empresa a un proveedor
            doc = rnd.choice(empresas)
            monto = float(rnd.randrange(200_000, 2_500_000, 10_000))
            if monto <= saldo[doc]:
                saldo[doc] -= monto
                agregar(fecha, "PAGO", doc, None, monto, "Pago a proveedor")
                continue
        if opcion < 0.85:                      # transferencia entre usuarios
            origen = rnd.choice(personas)
            destino = rnd.choice([d for d in todos if d != origen])
            tope = min(saldo[origen] * 0.2, 800_000)
            if tope >= 10_000:
                monto = round(rnd.uniform(10_000, tope), 2) if rnd.random() < 0.3 else \
                    float(rnd.randrange(10_000, int(tope) + 1, 1_000))
                saldo[origen] -= monto
                saldo[destino] += monto
                agregar(fecha, "TRANSFERENCIA", origen, destino, monto,
                        rnd.choice(["Pago de arriendo", "Prestamo", "Devolucion de prestamo",
                                    "Regalo", "Cuota compartida", "Pago de servicio"]))
                continue
        doc = rnd.choice(personas)              # recarga
        monto = float(rnd.randrange(50_000, 1_500_000, 10_000))
        saldo[doc] += monto
        agregar(fecha, "RECARGA", doc, None, monto, rnd.choice(["Recarga en corresponsal",
                                                                   "Recarga por PSE", "Recarga"]))

    # 4) Nomina
    for i, (nit, doc, salario) in enumerate(nomina):
        saldo[nit] -= salario
        saldo[doc] += salario
        agregar(datetime(2026, 9, 30, 16, 0) + timedelta(minutes=2 * i), "TRANSFERENCIA",
                nit, doc, salario, "Pago de nomina")

    assert len(movimientos) == TOTAL_MOVIMIENTOS, len(movimientos)
    assert all(v >= -0.001 for v in saldo.values())
    return movimientos


def escribir_hoja(libro, titulo, columnas, filas, formatos):
    hoja = libro.create_sheet(titulo)
    hoja.append(columnas)
    for fila in filas:
        hoja.append([fila[c] for c in columnas])
    encabezado = PatternFill("solid", fgColor="15304A")
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = encabezado
        celda.alignment = Alignment(horizontal="center", vertical="center")
    for i, columna in enumerate(columnas, start=1):
        letra = get_column_letter(i)
        ancho = max(len(columna), *(len(str(f[columna] or "")) for f in filas)) + 3
        hoja.column_dimensions[letra].width = min(ancho, 34)
        if columna in formatos:
            for celda in hoja[letra][1:]:
                if not isinstance(celda.value, str):
                    celda.number_format = formatos[columna]
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions
    return hoja


if __name__ == "__main__":
    rnd = random.Random(SEMILLA)
    usuarios, validos = generar_usuarios(rnd)
    movimientos = generar_movimientos(rnd, validos)

    libro = Workbook()
    libro.remove(libro.active)
    escribir_hoja(libro, "usuarios",
                  ["tipo_usuario", "nombre", "email", "cedula", "nit", "cargo", "salario",
                   "nit_empresa"],
                  usuarios, {"salario": "#,##0.00"})
    escribir_hoja(libro, "movimientos",
                  ["fecha", "tipo", "documento_origen", "documento_destino", "monto",
                   "descripcion"],
                  movimientos, {"fecha": "yyyy-mm-dd hh:mm", "monto": "#,##0.00"})
    libro.save(RUTA_EXCEL)

    print("=" * 66)
    print(" QUANTUM CORE - Generacion de datos simulados para la carga masiva")
    print("=" * 66)
    print(f"Archivo creado: {RUTA_EXCEL.name}")
    print(f"  Hoja usuarios    : {len(usuarios):>4} filas")
    print(f"  Hoja movimientos : {len(movimientos):>4} filas")
    print(f"  Total            : {len(usuarios) + len(movimientos):>4} registros")
